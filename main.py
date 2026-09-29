import cv2
import re
import numpy as np
import os
import sys
import json
import logging
import argparse
import threading

logging.disable(logging.WARNING)
from rapidocr import RapidOCR

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# -----------------------------------------------------------------------
# defaults & constants
# -----------------------------------------------------------------------

DEFAULTS = {
    "box_thresh": 0.5,
    "preprocess_mode": "无",
    "export_format": "Markdown",
    "show_boxes": True,
    "layout_analysis": True,
    "layout_conf_thresh": 0.3,
    "parallel": 0,
    "formula_ocr": False,
}

FMT_MAP = {"Markdown": "md", "Word": "docx", "TXT": "txt", "JSON": "json", "Excel": "xlsx"}


# -----------------------------------------------------------------------
# shared helpers (CLI / BmScriptsBox)
# -----------------------------------------------------------------------


def _preprocess(img, mode: str):
    """Image preprocessing before OCR."""
    if mode == "自动":
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        std = float(gray.std())
        mean = float(gray.mean())
        if std < 40:
            clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
            return cv2.cvtColor(clahe.apply(gray), cv2.COLOR_GRAY2BGR)
        if mean < 60 or mean > 200:
            p2, p98 = np.percentile(gray, (2, 98))
            stretched = np.clip(
                (gray.astype(np.float32) - p2) * (255.0 / max(p98 - p2, 1.0)),
                0, 255).astype(np.uint8)
            return cv2.cvtColor(stretched, cv2.COLOR_GRAY2BGR)
        return img
    if mode == "无":
        return img
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    if mode == "灰度化":
        return cv2.cvtColor(gray, cv2.COLOR_GRAY2BGR)
    if mode == "二值化":
        _, bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        return cv2.cvtColor(bw, cv2.COLOR_GRAY2BGR)
    if mode == "降噪":
        return cv2.medianBlur(img, 3)
    if mode == "增强对比度":
        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
        enhanced = clahe.apply(gray)
        return cv2.cvtColor(enhanced, cv2.COLOR_GRAY2BGR)
    return img


def _get_version():
    try:
        with open(os.path.join(SCRIPT_DIR, "pyproject.toml"), encoding="utf-8") as f:
            for line in f:
                if line.startswith("version"):
                    return line.split('"')[1]
    except OSError:
        pass
    return "?"


def _print_banner(fmt, box_thresh, preprocess, layout_on=True):
    ver = _get_version()
    print("=" * 54, file=sys.stderr)
    print(f"  BM-OCR v{ver} — 离线 OCR 文字识别工具", file=sys.stderr)
    print(f"  输出格式: {fmt.upper():5s} | "
          f"检测阈值: {box_thresh:.2f} | "
          f"预处理: {preprocess} | "
          f"版面分析: {'开' if layout_on else '关'}", file=sys.stderr)
    print("=" * 54, file=sys.stderr)
    print(file=sys.stderr)


def _print_batch_warning():
    print("由于批量OCR运算量较大，运行时会占用较多CPU资源。\n"
          "为了保证您的电脑操作依然流畅，并防止识别进程意外中断，\n"
          "建议您提交任务后，先关闭其他非必需的大型应用。\n", file=sys.stderr)


# -----------------------------------------------------------------------
# one-shot warnings (engine init happens per worker thread)
# -----------------------------------------------------------------------

_WARNED = set()
_WARN_LOCK = threading.Lock()


def _warn_once(key, msg):
    """同一类初始化失败只提示一次（并行时每个工作线程都会尝试初始化）。"""
    with _WARN_LOCK:
        if key in _WARNED:
            return
        _WARNED.add(key)
    print(msg, file=sys.stderr)


# -----------------------------------------------------------------------
# OCR engine construction
# -----------------------------------------------------------------------

# 方向分类整页投票。PP-OCR 的 cls 是 0°/180° 二分类，对短文本行误判率不低
# （实测 6 张图 204 个文本行里，3 行被误判并真的旋转，占 1.5%）。一旦判错，
# RapidOCR 会执行 cv2.rotate(img, ROTATE_180) 把正立的图转倒，识别结果随之
# 崩坏 —— 漏字、丢行，甚至 300mg 变 30mg 这种数字错位。
# 整页真的倒转时，几乎每一行都会判 180°；个别的误判只会是零星几行。
# 用「被判 180° 的行占比」就能把这两种情况干净地分开。
_CLS_VOTE_MIN_LINES = 3     # 行数太少时投票无意义，保持库的原始行为
_CLS_VOTE_180_RATIO = 0.5   # 判 180° 的行占比低于此值 → 视为误判，整批不旋转


def _make_ocr():
    """构造 RapidOCR，并装上方向分类的整页投票。

    RapidOCR 默认 use_cls=True + cls_thresh=0.9，会把正立的短文本行判成
    180° 并真的旋转图片。这里包一层 cls_and_rotate：先按库的逻辑跑，再统计
    被判 180° 的行占比，占比过低就把整批图退回未旋转的原图。

    返回值的 img_list 才是喂给识别模块的东西，所以只需替换它；
    cls_res 里的原始判定保持不动，便于排查。
    """
    eng = RapidOCR()
    _orig_cls_and_rotate = eng.cls_and_rotate

    def _voting_cls_and_rotate(img_list, *args, **kwargs):
        rotated, cls_res = _orig_cls_and_rotate(img_list, *args, **kwargs)
        labels = [label for label, _ in (cls_res.cls_res or [])]
        if (len(labels) >= _CLS_VOTE_MIN_LINES
                and labels.count("180") / len(labels) < _CLS_VOTE_180_RATIO):
            return list(img_list), cls_res
        return rotated, cls_res

    eng.cls_and_rotate = _voting_cls_and_rotate
    return eng


# -----------------------------------------------------------------------
# AI models — 由不忙脚本盒子统一下载，脚本只问路径，不自己下
# -----------------------------------------------------------------------
# 脚本在 rc.toml 的 [[bmscriptsbox.models]] 里声明所需模型，盒子安装时下好，
# 运行时用 GET {api_base}/api/model/path?repo_id=... 换取本地文件夹路径。
# 这样模型不进脚本包、多脚本共用一份，也避免各库各自往 site-packages 里囤模型。
# 盒子外（CLI 调试）查不到路径时回退到库自带的按需下载，保证脚本仍可用。

MODEL_LAYOUT_REPO = "RapidAI/RapidLayout"
MODEL_LAYOUT_FILE = "onnx/pp_doc_layout/pp_doc_layoutv3.onnx"
MODEL_FORMULA_REPO = "RapidAI/RapidDoc"
MODEL_FORMULA_FILE = "formula/PP-FormulaNet_plus-M/pp_formulanet_plus_m.onnx"

_MODEL_PATH_CACHE = {}
_MODEL_PATH_LOCK = threading.Lock()


def _query_model_path(api_base, repo_id, rel_file):
    """问盒子要模型文件夹，拼出具体模型文件路径。拿不到返回 None。"""
    from urllib.error import HTTPError
    from urllib.parse import quote
    from urllib.request import urlopen

    url = f"{api_base}/api/model/path?repo_id={quote(repo_id)}"
    try:
        with urlopen(url, timeout=5) as resp:
            body = json.loads(resp.read().decode("utf-8"))
    except HTTPError as e:
        # 404 = 盒子没下过这个模型，按「未提供」处理，交给调用方回退
        if e.code == 404:
            return None
        raise

    if not body.get("success"):
        return None
    path = os.path.join(body["data"]["path"], *rel_file.split("/"))
    return path if os.path.exists(path) else None


def _resolve_model(api_base, repo_id, rel_file):
    """解析模型本地路径；结果按 repo 缓存（并行时多个线程只查一次）。

    返回 None 表示盒子没给（未安装 / 盒子外运行），调用方应回退。
    """
    key = (api_base or "", repo_id)
    with _MODEL_PATH_LOCK:
        if key in _MODEL_PATH_CACHE:
            return _MODEL_PATH_CACHE[key]

    path = None
    if api_base:
        try:
            path = _query_model_path(api_base, repo_id, rel_file)
        except Exception as e:
            _warn_once("model_api",
                       f"未取到盒子下发的模型路径（{repo_id}），"
                       f"改用库自带下载: {e}")

    with _MODEL_PATH_LOCK:
        _MODEL_PATH_CACHE[key] = path
    return path


# -----------------------------------------------------------------------
# layout analysis (rapid_layout)
# -----------------------------------------------------------------------


def _ensure_layout_analyzer(conf_thresh=0.3, model_path=None):
    """Return a LayoutAnalyzer (rapid_layout) or None if unavailable.

    model_path 为盒子下载好的 pp_doc_layoutv3.onnx 路径；为空则退回
    rapid_layout 自带的按需下载（仅下载 v3 这一个）。
    """
    try:
        from src.layout_analyzer import LayoutAnalyzer
        return LayoutAnalyzer(conf_thresh=conf_thresh, model_path=model_path)
    except Exception as e:
        _warn_once("layout", f"版面分析初始化失败: {e}")
        return None


# -----------------------------------------------------------------------
# exports
# -----------------------------------------------------------------------


def _is_table_sep_row(s):
    """判断是否为 Markdown 表格分隔行（如 | --- | --- |、|:---:|）。"""
    cells = [c.strip() for c in s.strip().split("|")[1:-1]]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def _md_to_txt(md_text):
    plain = []
    for line in md_text.split("\n"):
        s = line.strip()
        if not s:
            plain.append("")
            continue
        if s.startswith("### "):
            plain.append(s[4:])
        elif s.startswith("## "):
            plain.append(s[3:])
        elif s.startswith("# "):
            plain.append(s[2:])
        elif s.startswith("|") and s.endswith("|"):
            if _is_table_sep_row(s):
                continue
            cells = [c.strip() for c in s.split("|")[1:-1]]
            plain.append(" | ".join(cells))
        elif s.startswith("- "):
            plain.append(s[2:])
        elif s.startswith("1. "):
            plain.append(s[3:])
        else:
            plain.append(s)
    return "\n".join(plain)


# -----------------------------------------------------------------------
# Word 公式：LaTeX → OMML（Word 原生公式）
# -----------------------------------------------------------------------

_OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _latex_to_omml(latex):
    """LaTeX → OMML 片段；依赖缺失或转换失败返回 None。

    走 latex2mathml → mathml2omml 两步，纯 Python，不需要本机装 Office。
    """
    try:
        import latex2mathml.converter as _l2m
        from mathml2omml import convert as _m2o
    except ImportError:
        return None
    try:
        omml = _m2o(_l2m.convert(latex))
    except Exception:
        return None
    if not omml:
        return None
    if not isinstance(omml, str):
        omml = str(omml)
    if "<m:oMath" not in omml:
        return None
    if "xmlns:m" not in omml:
        omml = omml.replace("<m:oMath", f'<m:oMath xmlns:m="{_OMML_NS}"', 1)
    return omml


def _add_docx_formula(doc, latex):
    """插入一个居中显示的原生 Word 公式段落。

    转换不可用时退化为普通段落保留 LaTeX 原文，避免公式内容丢失。
    Returns True 表示已插入原生公式。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    omml = _latex_to_omml(latex)
    if omml:
        try:
            p._p.append(parse_xml(
                f'<m:oMathPara xmlns:m="{_OMML_NS}">'
                f'<m:oMathParaPr><m:jc m:val="centerGroup"/></m:oMathParaPr>'
                f'{omml}</m:oMathPara>'))
            return True
        except Exception:
            pass
    p.add_run(latex)
    return False


def _md_to_docx(md_text, out_path):
    from docx import Document
    doc = Document()
    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("$$"):
            # 公式块：$$ 独占一行时收集到下一个 $$，同行闭合则直接取中间内容
            latex_lines = []
            if len(s) > 4 and s.endswith("$$"):
                latex_lines.append(s[2:-2].strip())
                i += 1
            else:
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("$$"):
                    latex_lines.append(lines[i].strip())
                    i += 1
                i += 1
            latex = " ".join(x for x in latex_lines if x).strip()
            if latex:
                _add_docx_formula(doc, latex)
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("|") and s.endswith("|"):
            rows = []
            while i < len(lines):
                sr = lines[i].strip()
                if not (sr.startswith("|") and sr.endswith("|")):
                    break
                if not _is_table_sep_row(sr):
                    rows.append(sr)
                i += 1
            if rows:
                ncols = rows[0].count("|") - 1
                tbl = doc.add_table(rows=len(rows), cols=ncols)
                tbl.style = "Table Grid"
                for ri, rstr in enumerate(rows):
                    cells = [c.strip() for c in rstr.split("|")[1:-1]]
                    for ci, ct in enumerate(cells):
                        if ci < ncols:
                            cell = tbl.rows[ri].cells[ci]
                            cell.text = ct
                            if ri == 0:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                doc.add_paragraph()
            continue
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif s.startswith("1. "):
            doc.add_paragraph(s[3:], style="List Number")
        else:
            doc.add_paragraph(s)
        i += 1
    doc.save(out_path)


def _export_excel(img_path, tables, md_text, out_path):
    """OCR 结果 → .xlsx：每个重建表格一个「表格N」sheet，非表格文本一个「文本」sheet。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def _write_table(ws, cells):
        """写入单元格并应用合并。

        顺序很关键：先写全部值、最后统一合并。openpyxl 的 MergedCell 只读
        （__slots__ 且 value 无 setter），先合并再往合并区内写值会抛错。

        另外 TableReconstructor 的 cell_bboxes 可能多于 HTML 解析出的单元格
        （rowspan/colspan 场景），多出来的"孤儿"单元格只有 box、没有 row/col，
        无法定位，这里跳过，而不是把它们全部堆到 A1 互相覆盖。
        """
        bold = Font(bold=True)
        widths = {}
        pending_merges = []

        for c in cells:
            row, col = c.get("row"), c.get("col")
            if row is None or col is None:
                continue
            try:
                row, col = int(row) + 1, int(col) + 1
            except (TypeError, ValueError):
                continue
            if row < 1 or col < 1:
                continue
            try:
                colspan = max(1, int(c.get("colspan") or 1))
                rowspan = max(1, int(c.get("rowspan") or 1))
            except (TypeError, ValueError):
                colspan = rowspan = 1

            text = c.get("text") or ""
            cell = ws.cell(row=row, column=col, value=text)
            if row == 1:
                cell.font = bold
            widths[col] = max(widths.get(col, 0), len(text))
            if colspan > 1 or rowspan > 1:
                pending_merges.append(
                    (row, col, row + rowspan - 1, col + colspan - 1))

        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = min(max(w * 1.8, 8), 50)

        merged = []
        for r1, c1, r2, c2 in pending_merges:
            if any(not (r2 < mr1 or r1 > mr2 or c2 < mc1 or c1 > mc2)
                   for mr1, mc1, mr2, mc2 in merged):
                continue
            try:
                ws.merge_cells(start_row=r1, start_column=c1,
                               end_row=r2, end_column=c2)
            except Exception:
                continue
            merged.append((r1, c1, r2, c2))

    table_sheets = 0
    for i, t in enumerate(tables or []):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = f"表格{i + 1}"
        _write_table(ws, t.get("cells", []))
        table_sheets += 1

    # 无重建表格时回退：解析 markdown 表格
    if not tables:
        md_rows = []
        for line in md_text.split("\n"):
            s = line.strip()
            if s.startswith("|") and s.endswith("|"):
                if _is_table_sep_row(s):
                    continue
                md_rows.append([c.strip() for c in s.split("|")[1:-1]])
        if md_rows:
            ws = wb.active if table_sheets == 0 else wb.create_sheet()
            ws.title = "表格1"
            cells = [{"row": ri, "col": ci, "text": text}
                     for ri, r in enumerate(md_rows) for ci, text in enumerate(r)]
            _write_table(ws, cells)
            table_sheets += 1

    # 文本 sheet：非表格 OCR 文本
    text_lines = []
    for line in md_text.split("\n"):
        s = line.strip()
        if not s or (s.startswith("|") and s.endswith("|")):
            continue
        if s.startswith("### "):
            text_lines.append(s[4:])
        elif s.startswith("## "):
            text_lines.append(s[3:])
        elif s.startswith("# "):
            text_lines.append(s[2:])
        elif s.startswith("- "):
            text_lines.append(s[2:])
        else:
            text_lines.append(s)
    if text_lines:
        ws = wb.active if table_sheets == 0 else wb.create_sheet()
        ws.title = "文本"
        for i, line in enumerate(text_lines, start=1):
            ws.cell(row=i, column=1, value=line)

    wb.save(out_path)


def _json_safe(obj):
    """Recursively convert numpy scalars/arrays to JSON-serializable types."""
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, np.integer):
        return int(obj)
    if isinstance(obj, np.floating):
        return float(obj)
    if isinstance(obj, (list, tuple)):
        return [_json_safe(x) for x in obj]
    if isinstance(obj, dict):
        return {k: _json_safe(v) for k, v in obj.items()}
    return obj


def _export_json(img_path, result, regions, out_path):
    boxes_list = [b.tolist() if hasattr(b, "tolist") else b for b in result.boxes]
    data = {
        "image": os.path.basename(img_path),
        "ocr": {
            "txts": list(result.txts),
            "scores": [float(s) for s in result.scores],
            "boxes": boxes_list,
            "elapse": _json_safe(result.elapse),
            "elapse_list": _json_safe(result.elapse_list) if result.elapse_list else [],
        },
        "layout_regions": _json_safe(regions) if regions else [],
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _draw_boxes(img, boxes):
    vis = img.copy()
    for box in boxes:
        b = np.asarray(box)
        ys, xs = b[:, 1], b[:, 0]
        x1, y1, x2, y2 = int(min(xs)), int(min(ys)), int(max(xs)), int(max(ys))
        cv2.rectangle(vis, (x1, y1), (x2, y2), (0, 255, 0), 2)
    return vis


def _build_ocr_results(result):
    """Per-line OCR results with coordinates: [{text, score, box, points}, ...]."""
    boxes = result.boxes if result.boxes is not None else []
    txts = result.txts if result.txts is not None else []
    scores = result.scores if result.scores is not None else []
    items = []
    for box, txt, score in zip(boxes, txts, scores):
        b = np.asarray(box)
        xs, ys = b[:, 0], b[:, 1]
        items.append({
            "text": txt,
            "score": float(score),
            "box": [int(min(xs)), int(min(ys)), int(max(xs)), int(max(ys))],
            "points": [[float(px), float(py)] for px, py in b.tolist()],
        })
    return items


# -----------------------------------------------------------------------
# per-image OCR
# -----------------------------------------------------------------------


def _count_boxes_in_bbox(boxes, bbox):
    """统计 OCR 文本框中心落在 region bbox 内的数量。"""
    if boxes is None:
        return 0
    x1, y1, x2, y2 = bbox
    n = 0
    for b in boxes:
        bb = np.asarray(b)
        if bb.size == 0 or bb.ndim != 2 or bb.shape[1] != 2:
            continue
        xs, ys = bb[:, 0], bb[:, 1]
        cx = (float(xs.min()) + float(xs.max())) / 2
        cy = (float(ys.min()) + float(ys.max())) / 2
        if x1 <= cx <= x2 and y1 <= cy <= y2:
            n += 1
    return n


def _reconstructed_cell_count(md_table):
    """统计重建表格的单元格数（rows × cols，不含分隔行）。"""
    data_rows = [l for l in md_table.split("\n")
                 if l.strip().startswith("|") and not l.strip().startswith("|---")]
    if not data_rows:
        return 0
    ncols = data_rows[0].count("|") - 1
    return len(data_rows) * max(1, ncols)


def _extract_region_ocr(result, bbox, ox, oy):
    """从全图 OCR 结果中提取落在 region 内的文本框，平移到 crop 坐标系。"""
    if result.boxes is None or len(result.boxes) == 0:
        return None
    x1, y1, x2, y2 = bbox
    boxes, txts, scores = [], [], []
    for b, t, s in zip(result.boxes, result.txts, result.scores):
        bb = np.asarray(b)
        if bb.size == 0 or bb.ndim != 2 or bb.shape[1] != 2:
            continue
        xs, ys = bb[:, 0], bb[:, 1]
        cx = (float(xs.min()) + float(xs.max())) / 2
        cy = (float(ys.min()) + float(ys.max())) / 2
        if x1 - 6 <= cx <= x2 + 6 and y1 - 6 <= cy <= y2 + 6:
            boxes.append(bb - [ox, oy])
            txts.append(t)
            scores.append(s)
    if not boxes:
        return None
    return np.array(boxes), txts, scores


def _reconstruct_tables(img, regions, result, engine, box_thresh, table_engine, max_dim=1000):
    """对版面检出的 table 区域做表格重建。

    Returns (table_markdowns, tables):
      table_markdowns: {(x1,y1,x2,y2): (md_table, cells)}
      tables: [{bbox, cells}] 用于 ocr_data，坐标在整图坐标系

    小表复用全图 OCR 结果（不重复推理）；大表缩放后同坐标系重建。
    质量门控：重建单元格数远少于 region 内 OCR 文本框数时判定失败，回退逐行排版。
    """
    if table_engine is None:
        return {}, []
    from src.table_reconstructor import crop_table_region

    table_markdowns = {}
    tables = []
    for r in regions:
        if r["label"] != "table":
            continue
        crop, (ox, oy) = crop_table_region(img, r["bbox"])
        if crop is None:
            continue
        h, w = crop.shape[:2]
        ocr_result = None
        if max(h, w) <= max_dim:
            ocr_result = _extract_region_ocr(result, r["bbox"], ox, oy)
        md_table, cells = table_engine.reconstruct(
            crop, engine, ocr_result=ocr_result, box_thresh=box_thresh)
        if md_table is None:
            continue

        region_boxes = _count_boxes_in_bbox(result.boxes, r["bbox"])
        cell_count = _reconstructed_cell_count(md_table)
        if region_boxes >= 4 and cell_count < max(2, int(region_boxes * 0.6)):
            print(f"  -> 表格重建不完整 ({cell_count} 单元格 / {region_boxes} 文本框)，回退逐行排版",
                  file=sys.stderr)
            continue

        key = tuple(round(float(v), 1) for v in r["bbox"])
        table_markdowns[key] = (md_table, cells)
        # 偏移回整图坐标系
        full_cells = []
        for c in (cells or []):
            x1, y1, x2, y2 = c["box"]
            full_cells.append({
                "row": c.get("row"),
                "col": c.get("col"),
                "colspan": c.get("colspan", 1),
                "rowspan": c.get("rowspan", 1),
                "text": c.get("text"),
                "box": [x1 + ox, y1 + oy, x2 + ox, y2 + oy],
            })
        tables.append({
            "bbox": [round(v, 1) for v in r["bbox"]],
            "cells": full_cells,
        })
    return table_markdowns, tables


def _ocr_image(img_path, engine, fmt, box_thresh, preprocess, show_boxes, analyzer,
               table_engine=None, formula_engine=None):
    """OCR a single image; write result alongside it.

    Returns (out_path, md_text, vis_path, ocr_items, tables, formulas).
    """
    if not os.path.exists(img_path):
        print(f"跳过: {img_path} (文件不存在)", file=sys.stderr)
        return None, None, None, None, None, None

    out_stem = os.path.splitext(os.path.basename(img_path))[0] + "_ocr"
    out_dir = os.path.dirname(img_path)
    print(f"处理: {img_path}", file=sys.stderr)

    try:
        raw = np.fromfile(img_path, dtype=np.uint8)
        img = cv2.imdecode(raw, cv2.IMREAD_COLOR)
        if img is None:
            print(f"  -> 无法读取图片", file=sys.stderr)
            return None, None, None, None, None, None

        if preprocess != "无":
            img = _preprocess(img, preprocess)

        result = engine(img, box_thresh=box_thresh,
                        return_word_box=True, return_single_char_box=True)

        regions = analyzer.analyze(img) if analyzer else None

        table_markdowns, tables = _reconstruct_tables(
            img, regions or [], result, engine, box_thresh, table_engine)
        formula_markdowns, formulas = _recognize_formulas(
            img, regions or [], formula_engine)

        from src.markdown_converter import MarkdownConverter
        converter = MarkdownConverter(
            result.boxes, result.txts, regions=regions,
            word_results=result.word_results,
            table_markdowns=table_markdowns,
            formula_markdowns=formula_markdowns,
        )
        md_text = converter.convert() if fmt != "json" else ""

        out_path = os.path.join(out_dir, f"{out_stem}.{fmt}")
        if fmt == "md":
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(md_text)
        elif fmt == "txt":
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(_md_to_txt(md_text))
        elif fmt == "docx":
            _md_to_docx(md_text, out_path)
        elif fmt == "json":
            _export_json(img_path, result, regions, out_path)
        elif fmt == "xlsx":
            _export_excel(img_path, tables, md_text, out_path)
        else:
            return None, None, None, None, None, None
        print(f"  -> {out_path}", file=sys.stderr)

        vis_path = None
        if show_boxes and result.boxes is not None and len(result.boxes):
            vis_path = os.path.join(out_dir, f"{out_stem}_boxes.png")
            vis = _draw_boxes(img, result.boxes)
            cv2.imencode(".png", vis)[1].tofile(vis_path)
            print(f"  -> {vis_path}", file=sys.stderr)

        return out_path, md_text, vis_path, _build_ocr_results(result), tables, formulas

    except Exception as e:
        print(f"  -> 出错: {e}", file=sys.stderr)
        return None, None, None, None, None, None


# =======================================================================
# BmScriptsBox entry (params_form / schedule / node)
# =======================================================================


class _EnginePool:
    """按线程持有 {ocr, analyzer, table, formula} 一套引擎。

    底层是 ONNXRuntime 会话，共享实例并发调用存在竞态，故每线程一份；
    且只在真正被该线程使用时构建 —— RapidOCR 构造即加载模型，
    这样批量 N 张图最多加载 min(线程数, N) 次，而不是 N 次。
    """

    def __init__(self, build):
        self._build = build
        self._local = threading.local()

    def get(self):
        ctx = getattr(self._local, "ctx", None)
        if ctx is None:
            ctx = self._build()
            self._local.ctx = ctx
        return ctx


def _run_images(files, fmt, box_thresh, preprocess, show_boxes,
                build_engines, workers=0):
    """批量 OCR。workers=0 时自动（最多 4 个线程），1 为串行。

    build_engines: 无参工厂，返回 {"ocr","analyzer","table","formula"}，
                   由每个工作线程各自构建一份。
    Returns list of _ocr_image results，顺序与 files 一致。
    """
    if not files:
        return []

    if workers and workers > 0:
        n = int(workers)
    else:
        n = min(os.cpu_count() or 2, 4)
    n = max(1, min(n, 8, len(files)))

    pool = _EnginePool(build_engines)

    def _run_one(p):
        ctx = pool.get()
        return _ocr_image(p, ctx["ocr"], fmt, box_thresh, preprocess, show_boxes,
                          ctx["analyzer"], ctx["table"], ctx["formula"])

    if n == 1:
        return [_run_one(p) for p in files]

    from concurrent.futures import ThreadPoolExecutor

    ordered = {}
    with ThreadPoolExecutor(max_workers=n) as ex:
        for idx, r in ex.map(lambda a: (a[0], _run_one(a[1])), enumerate(files)):
            ordered[idx] = r
    return [ordered[i] for i in range(len(files))]


def run_bmscripts(params_path):
    """Box mode: python main.py <path-to-params.json>

    Reads target_paths from data and run options from params (TOML defaults
    merged by the box). Writes an envelope when invoked as a node or when
    output_json is provided (scheduled-task reporting).
    """
    with open(params_path, "r", encoding="utf-8") as f:
        payload = json.load(f)

    env = payload.get("environment", {})
    invoke_mode = env.get("invoke_mode", "manual")
    params = payload.get("params", {})
    files = payload.get("data", {}).get("target_paths", []) or []

    box_thresh = float(params.get("box_thresh", DEFAULTS["box_thresh"]))
    preprocess_mode = params.get("preprocess_mode", DEFAULTS["preprocess_mode"])
    export_format = params.get("export_format", DEFAULTS["export_format"])
    show_boxes = bool(params.get("show_boxes", DEFAULTS["show_boxes"]))
    layout_analysis = bool(params.get("layout_analysis", DEFAULTS["layout_analysis"]))
    layout_conf_thresh = float(params.get("layout_conf_thresh", DEFAULTS["layout_conf_thresh"]))
    parallel = int(params.get("parallel", DEFAULTS["parallel"]))
    formula_ocr = bool(params.get("formula_ocr", DEFAULTS["formula_ocr"]))
    fmt = FMT_MAP.get(export_format, "md")

    # 模型路径一次问清（并行时多线程共用），拿不到就留给各库自行按需下载
    api_base = env.get("api_base")
    layout_model_path = (_resolve_model(api_base, MODEL_LAYOUT_REPO, MODEL_LAYOUT_FILE)
                         if layout_analysis else None)
    formula_model_path = (_resolve_model(api_base, MODEL_FORMULA_REPO, MODEL_FORMULA_FILE)
                          if formula_ocr else None)

    _print_banner(fmt, box_thresh, preprocess_mode, layout_analysis)

    def _build_engines():
        """每个工作线程各自构建一套引擎（ONNX 会话不可跨线程共享）。"""
        return {
            "ocr": _make_ocr(),
            "analyzer": (_ensure_layout_analyzer(layout_conf_thresh, layout_model_path)
                         if layout_analysis else None),
            "table": _ensure_table_engine(),
            "formula": (_ensure_formula_engine(formula_model_path)
                        if formula_ocr else None),
        }

    if len(files) > 1:
        _print_batch_warning()

    results = []
    texts = []
    ocr_data = []
    ok_count = 0
    for p, (out_path, md_text, _, items, tables, formulas) in zip(
            files,
            _run_images(files, fmt, box_thresh, preprocess_mode, show_boxes,
                        _build_engines, parallel)):
        if out_path:
            ok_count += 1
            results.append(out_path)
        if md_text:
            texts.append(md_text)
        if items:
            entry = {"image": os.path.basename(p), "items": items}
            if tables:
                entry["tables"] = tables
            if formulas:
                entry["formulas"] = formulas
            ocr_data.append(entry)

    text = "\n\n".join(texts)
    msg = f"已识别 {ok_count}/{len(files)} 张图片" if files else "没有待处理的图片"

    if invoke_mode == "node":
        code = 0 if ok_count > 0 else 1
        with open(env["output_json"], "w", encoding="utf-8") as f:
            json.dump({"code": code, "msg": msg, "results": results,
                       "text": text, "ocr_data": ocr_data},
                      f, ensure_ascii=False)
    elif env.get("output_json"):
        code = 0 if ok_count > 0 else 1
        with open(env["output_json"], "w", encoding="utf-8") as f:
            json.dump({"code": code, "msg": msg}, f, ensure_ascii=False)
    else:
        print(msg, file=sys.stderr)

    return 0 if ok_count > 0 else 1


# =======================================================================
# CLI (dev / standalone)
# =======================================================================


def _ensure_table_engine():
    """Return a TableReconstructor (rapid_table, SLANetPlus bundled) or None."""
    try:
        from src.table_reconstructor import TableReconstructor
        return TableReconstructor()
    except Exception as e:
        _warn_once("table", f"表格重建初始化失败: {e}")
        return None


def _ensure_formula_engine(model_path=None):
    """Return a FormulaRecognizer (PP-FormulaNet_plus-M) or None.

    model_path 为盒子下载好的 pp_formulanet_plus_m.onnx 路径；为空则退回
    库自带的按需下载（仅下载 M 这一个）。
    """
    try:
        from src.formula_recognizer import FormulaRecognizer
        return FormulaRecognizer(model_path=model_path)
    except Exception as e:
        _warn_once("formula", f"公式识别初始化失败: {e}")
        return None


def _recognize_formulas(img, regions, formula_engine):
    """对版面检出的 equation 区域做公式识别。

    Returns (formula_markdowns, formulas):
      formula_markdowns: {(x1,y1,x2,y2): latex}
      formulas: [{bbox, latex}] 用于 ocr_data
    """
    if formula_engine is None:
        return {}, []
    from src.formula_recognizer import crop_region

    formula_markdowns = {}
    formulas = []
    for r in regions:
        if r["label"] != "equation":
            continue
        crop = crop_region(img, r["bbox"])
        if crop is None:
            continue
        latex = formula_engine.recognize(crop)
        if not latex:
            continue
        key = tuple(round(float(v), 1) for v in r["bbox"])
        formula_markdowns[key] = latex
        formulas.append({"bbox": [round(v, 1) for v in r["bbox"]], "latex": latex})
    return formula_markdowns, formulas


def run_cli(opts):
    fmt = "md" if opts.format in ("md", "markdown") else opts.format
    _print_banner(fmt, opts.box_thresh, opts.preprocess, opts.layout_analysis)

    # CLI 一般跑在盒子外：BM_API_BASE 未设时模型路径为 None，
    # 各识别库会退回自带的按需下载（也只下当前用的这一个模型）。
    api_base = os.environ.get("BM_API_BASE")
    layout_model_path = (_resolve_model(api_base, MODEL_LAYOUT_REPO, MODEL_LAYOUT_FILE)
                         if opts.layout_analysis else None)
    formula_model_path = (_resolve_model(api_base, MODEL_FORMULA_REPO, MODEL_FORMULA_FILE)
                          if opts.formula_ocr else None)

    def _build_engines():
        """每个工作线程各自构建一套引擎（ONNX 会话不可跨线程共享）。"""
        return {
            "ocr": _make_ocr(),
            "analyzer": (_ensure_layout_analyzer(opts.layout_conf_thresh, layout_model_path)
                         if opts.layout_analysis else None),
            "table": _ensure_table_engine(),
            "formula": (_ensure_formula_engine(formula_model_path)
                        if opts.formula_ocr else None),
        }

    if len(opts.images) > 1:
        _print_batch_warning()

    _run_images(opts.images, fmt, opts.box_thresh, opts.preprocess,
                opts.show_boxes, _build_engines, opts.parallel)
    return 0


def _pause_cli(timeout=5):
    """等回车或超时后退出。

    不能另起线程去阻塞 input()：线程是 daemon，解释器退出时它仍持有 stdin 锁，
    会抛 "_enter_buffered_busy: could not acquire lock" 致命错误。这里改为
    主线程轮询，超时即返回。
    """
    print(f"\n按回车键退出 ({timeout}秒后自动退出)...", end="", flush=True)
    try:
        if os.name == "nt":
            import msvcrt
            import time as _time
            deadline = _time.monotonic() + timeout
            while _time.monotonic() < deadline:
                if msvcrt.kbhit():
                    msvcrt.getwch()
                    return
                _time.sleep(0.05)
        else:
            import select
            ready, _, _ = select.select([sys.stdin], [], [], timeout)
            if ready:
                sys.stdin.readline()
    except Exception:
        # 非交互环境（管道/重定向）直接跳过等待
        pass


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1].endswith(".json"):
        sys.exit(run_bmscripts(sys.argv[1]))

    elif len(sys.argv) >= 2:
        parser = argparse.ArgumentParser(description="离线OCR批处理工具")
        parser.add_argument("images", nargs="+", help="图片文件路径")
        parser.add_argument("--format", "-f",
                            choices=["txt", "json", "md", "markdown", "docx", "xlsx"],
                            default="md", help="输出格式")
        parser.add_argument("--box-thresh", type=float,
                            default=DEFAULTS["box_thresh"], help="检测阈值 (0.1-0.9)")
        parser.add_argument("--preprocess",
                            choices=["无", "自动", "灰度化", "二值化", "降噪", "增强对比度"],
                            default=DEFAULTS["preprocess_mode"], help="预处理模式")
        parser.add_argument("--show-boxes", action=argparse.BooleanOptionalAction,
                            default=DEFAULTS["show_boxes"], help="是否保存检测框标注图")
        parser.add_argument("--layout-analysis", action=argparse.BooleanOptionalAction,
                            default=DEFAULTS["layout_analysis"], help="是否启用版面分析")
        parser.add_argument("--layout-conf-thresh", type=float,
                            default=DEFAULTS["layout_conf_thresh"],
                            help="版面区域检测置信度阈值 (0.1-0.9)")
        parser.add_argument("--parallel", type=int,
                            default=DEFAULTS["parallel"],
                            help="批量并行数，0=自动(按CPU)，1=串行")
        parser.add_argument("--formula-ocr", action=argparse.BooleanOptionalAction,
                            default=DEFAULTS["formula_ocr"],
                            help="是否启用公式识别（LaTeX，较慢）")
        opts = parser.parse_args()
        rc = run_cli(opts)
        _pause_cli()
        sys.exit(rc)

    else:
        print("BM-OCR 离线文字识别工具（不忙脚本盒子脚本）", file=sys.stderr)
        print("用法:", file=sys.stderr)
        print("  盒子模式: python main.py <参数JSON路径>", file=sys.stderr)
        print("  CLI 调试: python main.py <图片...> [-f md|txt|docx|json|xlsx] "
              "[--box-thresh N] [--preprocess MODE] [--show-boxes] "
              "[--layout-analysis/--no-layout-analysis] "
              "[--layout-conf-thresh N] [--parallel N] [--formula-ocr]",
              file=sys.stderr)
        sys.exit(0)